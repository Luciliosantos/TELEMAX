"""Painel administrativo: proteção por role, moderação e configurações."""
import os
import math
import csv
import io
import json
import re
import unicodedata

from flask import (Blueprint, abort, flash, g, redirect, render_template,
                   request, url_for)

from .db import query, execute, get_db
from .security import (audit, csrf_ok, hash_password, new_token, now, save_image,
                       slugify, unique_slug, valid_telegram_link, normalize_link)

bp = Blueprint("admin", __name__)

@bp.before_request
def require_admin():
    u = g.get("user")
    if not u:
        return redirect(url_for("auth.login", next=request.path))
    if u["role"] != "admin":
        abort(403)

@bp.route("/admin")
def dashboard():
    stats = {
        "usuarios": query("SELECT COUNT(*) n FROM users", one=True)["n"],
        "grupos_pendentes": query("SELECT COUNT(*) n FROM telegram_groups WHERE status='pending'", one=True)["n"],
        "grupos_aprovados": query("SELECT COUNT(*) n FROM telegram_groups WHERE status='approved'", one=True)["n"],
        "promocoes_ativas": query("SELECT COUNT(*) n FROM promotions WHERE is_active=1 AND ends_at>?", (now(),), one=True)["n"],
        "receita_pago": query("SELECT COALESCE(SUM(amount_cents),0) n FROM payments WHERE status='paid' AND COALESCE(provider, '') != 'demo'", one=True)["n"],
        "pagamentos_pendentes": query("SELECT COUNT(*) n FROM payments WHERE status='pending'", one=True)["n"],
    }
    recent = query("SELECT pay.*, u.nickname, tg.name AS group_name FROM payments pay "
                   "LEFT JOIN users u ON u.id = pay.user_id "
                   "LEFT JOIN telegram_groups tg ON tg.id = pay.group_id "
                   "ORDER BY pay.created_at DESC LIMIT 10")
    logs = query("SELECT a.*, u.nickname FROM audit_log a LEFT JOIN users u ON u.id=a.actor_id "
                 "ORDER BY a.id DESC LIMIT 15")
    contatos = query("SELECT COUNT(*) n FROM contact_messages WHERE handled=0", one=True)["n"]
    return render_template("admin/dashboard.html", stats=stats, recent=recent, logs=logs,
                           contatos=contatos)

@bp.route("/admin/grupos")
def groups_admin():
    st = request.args.get("status", "pending")
    page = request.args.get("page", 1, type=int)
    per = 20
    total = query("SELECT COUNT(*) n FROM telegram_groups WHERE status = ?", (st,), one=True)["n"]
    rows = query("SELECT tg.*, c.name AS category_name, u.nickname AS owner "
                 "FROM telegram_groups tg LEFT JOIN categories c ON c.id=tg.category_id "
                 "LEFT JOIN users u ON u.id=tg.owner_id WHERE tg.status = ? "
                 "ORDER BY tg.created_at DESC LIMIT ? OFFSET ?",
                 (st, per, (page-1)*per))
    return render_template("admin/grupos.html", rows=rows, st=st, page=page,
                           pages=math.ceil(total/per) if total else 0)

@bp.route("/admin/grupos/<int:gid>/decidir", methods=["POST"])
def moderate(gid):
    if not csrf_ok():
        return "CSRF inválido", 400
    acao = request.form.get("acao")
    motivo = (request.form.get("motivo") or "").strip()[:300]
    row = query("SELECT * FROM telegram_groups WHERE id = ?", (gid,), one=True)
    if not row:
        abort(404)
    if acao == "aprovar":
        execute("UPDATE telegram_groups SET status='approved', published_at=?, rejection_reason=NULL "
                "WHERE id = ?", (now(), gid))
        flash("Publicação aprovada.", "ok")
    elif acao == "rejeitar":
        execute("UPDATE telegram_groups SET status='rejected', rejection_reason=? WHERE id = ?",
                (motivo or "Não atende às regras.", gid))
        flash("Publicação rejeitada.", "ok")
    elif acao == "remover":
        execute("UPDATE telegram_groups SET status='removed' WHERE id = ?", (gid,))
        flash("Publicação removida.", "ok")
    elif acao == "editar":
        execute("UPDATE telegram_groups SET name=?, description=?, link=?, keywords=? WHERE id=?",
                ((request.form.get("name") or row["name"]).strip()[:100],
                 (request.form.get("description") or row["description"]).strip()[:500],
                 (request.form.get("link") or row["link"]).strip(),
                 (request.form.get("keywords") or row["keywords"]).strip()[:200], gid))
        flash("Publicação editada.", "ok")
    audit(g.user["id"], f"group_{acao}", "group", gid, motivo, request.remote_addr)
    return redirect(request.referrer or url_for("admin.groups_admin"))


def _bulk_key(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", value).strip().casefold()


def _bulk_header_map(raw):
    aliases = {
        "nome": ("nome", "name", "grupo", "canal", "titulo", "title"),
        "link": ("link_telegram", "link", "url", "telegram", "link do telegram", "telegram_url"),
        "categoria": ("categoria", "category", "categoria do grupo", "categoria do canal"),
        "descricao": ("descricao", "descrição", "description", "detalhes", "sobre"),
        "keywords": ("keywords", "palavras_chave", "palavras-chave", "tags"),
        "tipo": ("tipo", "type", "tipo de grupo", "tipo do canal", "group_type"),
    }
    source = {_bulk_key(k): ("" if v is None else str(v).strip()) for k, v in raw.items() if k is not None}
    out = {}
    for key, names in aliases.items():
        out[key] = next((source[_bulk_key(n)] for n in names if _bulk_key(n) in source), "")
    return out


def _bulk_parse_text_rows(text, source_name="arquivo"):
    rows = []
    for line in text.splitlines():
        line = line.strip().strip("|")
        if not line or not re.search(r"(?:https?://)?(?:t\.me|telegram\.me|telegram\.dog)/", line, re.I):
            continue
        url_match = re.search(r"https?://(?:t\.me|telegram\.me|telegram\.dog)/[^\s|;,]+|(?:t\.me|telegram\.me|telegram\.dog)/[^\s|;,]+", line, re.I)
        if not url_match:
            continue
        url = url_match.group(0)
        if not url.lower().startswith("http"):
            url = "https://" + url
        parts = [x.strip() for x in re.split(r"\s*\|\s*|\s*;\s*|\t+|\s{2,}", line) if x.strip()]
        before = line[:url_match.start()].strip(" |;\t-")
        after = line[url_match.end():].strip(" |;\t-")
        if len(parts) >= 2:
            if parts[0].lower() in ("nome", "name", "grupo", "canal"):
                continue
            rows.append({"nome": parts[0], "link_telegram": url,
                         "categoria": parts[2] if len(parts) >= 3 else "Outros",
                         "descricao": parts[3] if len(parts) >= 4 else "",
                         "tipo": parts[4] if len(parts) >= 5 else ""})
        elif before:
            rows.append({"nome": before[:100], "link_telegram": url,
                         "categoria": "Outros", "descricao": after[:500], "tipo": ""})
    if not rows:
        raise ValueError(f"Não encontrei linhas com links do Telegram em {source_name}. No PDF, use texto selecionável e colunas separadas por |, ; ou tab. PDF escaneado precisa de OCR antes da importação.")
    return rows


def _bulk_parse_upload(filename, data):
    ext = os.path.splitext((filename or "").lower())[1]
    if len(data) > 5 * 1024 * 1024:
        raise ValueError("O arquivo deve ter no máximo 5 MB.")
    if ext == ".csv":
        try:
            content = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = data.decode("cp1252")
        reader = csv.DictReader(io.StringIO(content))
        if not reader.fieldnames:
            raise ValueError("CSV sem cabeçalho. Use nome, link_telegram, categoria e descricao.")
        return list(reader)
    if ext in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            values = workbook.active.iter_rows(values_only=True)
            headers = next(values, None)
            if not headers:
                raise ValueError("A planilha está vazia.")
            keys = [str(v).strip() if v is not None else "" for v in headers]
            rows = [dict(zip(keys, row)) for row in values if any(v is not None and str(v).strip() for v in row)]
            workbook.close()
            return rows
        except ImportError as exc:
            raise ValueError("Suporte a XLSX indisponível: instale openpyxl.") from exc
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Não foi possível ler a planilha XLSX.") from exc
    if ext == ".xls":
        try:
            import xlrd
            workbook = xlrd.open_workbook(file_contents=data)
            sheet = workbook.sheet_by_index(0)
            if not sheet.nrows:
                raise ValueError("A planilha está vazia.")
            headers = [str(v).strip() for v in sheet.row_values(0)]
            return [dict(zip(headers, sheet.row_values(i))) for i in range(1, sheet.nrows) if any(str(v).strip() for v in sheet.row_values(i))]
        except ImportError as exc:
            raise ValueError("Suporte a XLS indisponível: instale xlrd.") from exc
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Não foi possível ler a planilha XLS.") from exc
    if ext == ".json":
        try:
            obj = json.loads(data.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValueError("JSON inválido. Salve o arquivo em UTF-8.") from exc
        if isinstance(obj, dict):
            obj = obj.get("grupos", obj.get("items", obj.get("data", [obj])))
        if not isinstance(obj, list):
            raise ValueError("O JSON deve conter uma lista de grupos/canais ou a chave 'grupos'.")
        return [r for r in obj if isinstance(r, dict)]
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
        except ImportError as exc:
            raise ValueError("Suporte a PDF indisponível: instale pypdf.") from exc
        except Exception as exc:
            raise ValueError("Não foi possível ler o PDF. Verifique se está protegido ou corrompido.") from exc
        if not text.strip():
            raise ValueError("PDF escaneado/sem texto selecionável. Aplique OCR antes da importação.")
        return _bulk_parse_text_rows(text, "PDF")
    if ext == ".txt":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("TXT inválido. Salve em UTF-8.") from exc
        return _bulk_parse_text_rows(text, "TXT")
    raise ValueError("Formato não aceito. Envie CSV, XLSX, XLS, JSON, PDF ou TXT.")


def _bulk_prepare(raw_rows):
    categories = query("SELECT id, name FROM categories ORDER BY name")
    cat_map = {_bulk_key(c["name"]): c["id"] for c in categories}
    existing_links = {normalize_link(r["link"]) for r in query("SELECT link FROM telegram_groups")}
    seen = set()
    result = []
    categorias_novas_vistas = set()
    for number, raw in enumerate(raw_rows, start=2):
        if not isinstance(raw, dict):
            continue
        fields = _bulk_header_map(raw)
        name = fields["nome"][:100]
        link = normalize_link(fields["link"])
        category_name = (fields["categoria"] or "Outros").strip()[:80] or "Outros"
        description = fields["descricao"][:500]
        keywords = fields["keywords"][:200]
        raw_type = _bulk_key(fields["tipo"])
        group_type = "channel" if raw_type in ("canal", "channel", "broadcast", "canal telegram") else "group"
        category_id = cat_map.get(_bulk_key(category_name))
        reason = ""
        if not 3 <= len(name) <= 100:
            reason = "Nome deve ter entre 3 e 100 caracteres."
        elif not valid_telegram_link(link):
            reason = "Link do Telegram inválido. Use um link t.me público ou de convite."
        elif link in existing_links or link in seen:
            reason = "Link duplicado no site ou no próprio arquivo."
        if link:
            seen.add(link)
        category_key = _bulk_key(category_name)
        if not category_id and category_key:
            if category_key in categorias_novas_vistas:
                category_message = "Categoria nova já prevista nesta importação"
            else:
                category_message = "Categoria será criada automaticamente"
                categorias_novas_vistas.add(category_key)
        else:
            category_message = "Pronto para importar e publicar"
        result.append({"linha": number, "nome": name, "link": link, "categoria": category_name,
                       "categoria_id": category_id, "categoria_nova": not bool(category_id),
                       "descricao": description, "keywords": keywords, "tipo": group_type,
                       "valido": not bool(reason),
                       "motivo": reason or category_message})
    return result


def _bulk_get_or_create_category(db, category_name):
    name = (category_name or "Outros").strip()[:80] or "Outros"
    row = db.execute("SELECT id FROM categories WHERE lower(name)=lower(?) LIMIT 1", (name,)).fetchone()
    if row:
        return row[0], False
    base = slugify(name, "categoria")
    slug, suffix = base, 2
    while db.execute("SELECT id FROM categories WHERE slug=?", (slug,)).fetchone():
        slug, suffix = f"{base}-{suffix}", suffix + 1
    cur = db.execute("INSERT INTO categories(slug, name, description, created_at) VALUES (?,?,?,?)",
                     (slug, name, f"Grupos e canais sobre {name}.", now()))
    return cur.lastrowid, True


@bp.route("/admin/grupos/importar", methods=["GET", "POST"])
def bulk_import_groups():
    if request.method == "GET":
        return render_template("admin/importar_grupos.html", rows=None,
                               cats=query("SELECT * FROM categories ORDER BY name"))
    if not csrf_ok():
        return "CSRF inválido", 400
    action = request.form.get("action", "")
    if action == "preview":
        upload = request.files.get("arquivo")
        if not upload or not upload.filename:
            flash("Selecione um arquivo para importar.", "erro")
            return redirect(url_for("admin.bulk_import_groups"))
        data = upload.read(5 * 1024 * 1024 + 1)
        if len(data) > 5 * 1024 * 1024:
            flash("O arquivo deve ter no máximo 5 MB.", "erro")
            return redirect(url_for("admin.bulk_import_groups"))
        try:
            raw_rows = _bulk_parse_upload(upload.filename, data)
            if not raw_rows:
                raise ValueError("O arquivo não contém registros.")
            if len(raw_rows) > 500:
                raise ValueError("Limite de 500 registros por arquivo.")
            rows = _bulk_prepare(raw_rows)
        except (ValueError, csv.Error) as exc:
            flash("Não foi possível ler o arquivo: " + str(exc), "erro")
            return redirect(url_for("admin.bulk_import_groups"))
        return render_template("admin/importar_grupos.html", rows=rows,
                               cats=query("SELECT * FROM categories ORDER BY name"))
    if action == "confirm":
        try:
            raw_rows = json.loads(request.form.get("payload", ""))
            if not isinstance(raw_rows, list) or not raw_rows or len(raw_rows) > 500:
                raise ValueError("Prévia inválida ou vazia.")
            csv_rows = []
            for item in raw_rows:
                if not isinstance(item, dict):
                    raise ValueError("Registro da prévia inválido.")
                csv_rows.append({"nome": item.get("nome", ""), "link_telegram": item.get("link", ""),
                                 "categoria": item.get("categoria", "Outros"),
                                 "descricao": item.get("descricao", ""), "keywords": item.get("keywords", ""),
                                 "tipo": item.get("tipo", "group")})
            rows = _bulk_prepare(csv_rows)
        except (ValueError, TypeError, json.JSONDecodeError) as exc:
            flash("Prévia inválida. Envie o arquivo novamente. " + str(exc), "erro")
            return redirect(url_for("admin.bulk_import_groups"))
        valid_rows = [r for r in rows if r["valido"]]
        if not valid_rows:
            flash("Nenhum registro válido para importar. Confira os motivos na prévia.", "erro")
            return render_template("admin/importar_grupos.html", rows=rows,
                                   cats=query("SELECT * FROM categories ORDER BY name"))
        db = get_db()
        inserted, created_categories = 0, set()
        try:
            db.execute("BEGIN IMMEDIATE")
            existing_links = {
                normalize_link(r[0])
                for r in db.execute("SELECT link FROM telegram_groups").fetchall()
            }
            for row in valid_rows:
                if row["link"] in existing_links:
                    continue
                category_id, created = _bulk_get_or_create_category(db, row["categoria"])
                if created:
                    created_categories.add(row["categoria"])
                slug = unique_slug(slugify(row["nome"], "grupo"), "telegram_groups")
                db.execute("INSERT INTO telegram_groups "
                           "(slug, name, description, link, category_id, type, keywords, owner_id, status, is_demo, created_at, published_at) "
                           "VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'pending', 0, ?, NULL)",
                           (slug, row["nome"], row["descricao"], row["link"], category_id,
                            row["tipo"], row["keywords"], g.user["id"], now()))
                existing_links.add(row["link"])
                inserted += 1
            db.commit()
        except Exception:
            db.rollback()
            raise
        detail = f"Importação: {inserted} grupos/canais enviados para moderação; categorias criadas: {len(created_categories)}."
        audit(g.user["id"], "bulk_import", "group", None, detail, request.remote_addr)
        message = f"Importação concluída: {inserted} grupos/canais enviados para aprovação manual."
        if created_categories:
            message += " Categorias criadas: " + ", ".join(sorted(created_categories)) + "."
        if inserted < len(valid_rows):
            message += f" {len(valid_rows) - inserted} registro(s) foram ignorados por duplicidade detectada na confirmação."
        flash(message, "ok")
        return redirect(url_for("admin.groups_admin", status="pending"))
    return "Ação inválida", 400


@bp.route("/admin/usuarios")
def users_admin():
    page = request.args.get("page", 1, type=int)
    per = 20
    total = query("SELECT COUNT(*) n FROM users", one=True)["n"]
    rows = query("SELECT u.*, (SELECT COUNT(*) FROM telegram_groups tg WHERE tg.owner_id=u.id) AS grupos "
                 "FROM users u ORDER BY u.created_at DESC LIMIT ? OFFSET ?", (per, (page-1)*per))
    return render_template("admin/usuarios.html", rows=rows, page=page,
                           pages=math.ceil(total/per) if total else 0)

@bp.route("/admin/usuarios/<int:uid>/acao", methods=["POST"])
def user_action(uid):
    if not csrf_ok():
        return "CSRF inválido", 400
    acao = request.form.get("acao")
    if uid == g.user["id"] and acao != "resetar_sessoes":
        flash("Você não pode alterar sua própria conta por aqui.", "erro")
        return redirect(url_for("admin.users_admin"))
    if acao == "desativar":
        execute("UPDATE users SET is_active=0 WHERE id=?", (uid,))
        execute("UPDATE user_sessions SET revoked_at=? WHERE user_id=? AND revoked_at IS NULL", (now(), uid))
    elif acao == "ativar":
        execute("UPDATE users SET is_active=1, failed_logins=0, locked_until=NULL WHERE id=?", (uid,))
    elif acao == "tornar_admin":
        execute("UPDATE users SET role='admin' WHERE id=?", (uid,))
    elif acao == "remover_admin":
        execute("UPDATE users SET role='user' WHERE id=?", (uid,))
    audit(g.user["id"], f"user_{acao}", "user", uid, "", request.remote_addr)
    return redirect(url_for("admin.users_admin"))

@bp.route("/admin/categorias", methods=["GET", "POST"])
def categories_admin():
    if request.method == "POST":
        if not csrf_ok():
            return "CSRF inválido", 400
        name = (request.form.get("name") or "").strip()[:60]
        desc = (request.form.get("description") or "").strip()[:300]
        if len(name) >= 3:
            execute("INSERT INTO categories(slug, name, description, created_at) VALUES (?,?,?,?)",
                    (unique_slug(slugify(name, "cat"), "categories"), name, desc, now()))
            flash("Categoria criada.", "ok")
        return redirect(url_for("admin.categories_admin"))
    rows = query("SELECT c.*, (SELECT COUNT(*) FROM telegram_groups tg WHERE tg.category_id=c.id) AS total "
                 "FROM categories c ORDER BY c.name")
    return render_template("admin/categorias.html", rows=rows)

@bp.route("/admin/categorias/<int:cid>/acao", methods=["POST"])
def category_action(cid):
    if not csrf_ok():
        return "CSRF inválido", 400
    acao = request.form.get("acao")
    if acao == "renomear":
        name = (request.form.get("name") or "").strip()[:60]
        if len(name) >= 3:
            execute("UPDATE categories SET name=? WHERE id=?", (name, cid))
    elif acao == "excluir":
        n = query("SELECT COUNT(*) n FROM telegram_groups WHERE category_id=? AND status='approved'",
                  (cid,), one=True)["n"]
        if n == 0:
            execute("DELETE FROM categories WHERE id=?", (cid,))
            flash("Categoria excluída.", "ok")
        else:
            flash("Categoria possui publicações aprovadas.", "erro")
    audit(g.user["id"], f"category_{acao}", "category", cid, "", request.remote_addr)
    return redirect(url_for("admin.categories_admin"))

@bp.route("/admin/planos", methods=["GET", "POST"])
def plans_admin():
    if request.method == "POST":
        if not csrf_ok():
            return "CSRF inválido", 400
        name = (request.form.get("name") or "").strip()[:60]
        days = request.form.get("days", type=int)
        reais = request.form.get("price") or "0"
        try:
            cents = int(round(float(reais.replace(",", ".")) * 100))
        except ValueError:
            cents = -1
        desc = (request.form.get("description") or "").strip()[:200]
        if name and days and days > 0 and cents >= 0:
            execute("INSERT INTO plans(name, days, price_cents, description, sort, created_at) "
                    "VALUES (?,?,?,?,0,?)", (name, days, cents, desc, now()))
            flash("Plano criado.", "ok")
        else:
            flash("Dados do plano inválidos.", "erro")
        return redirect(url_for("admin.plans_admin"))
    rows = query("SELECT * FROM plans ORDER BY days")
    return render_template("admin/planos.html", rows=rows)

@bp.route("/admin/planos/<int:pid>/acao", methods=["POST"])
def plan_action(pid):
    if not csrf_ok():
        return "CSRF inválido", 400
    acao = request.form.get("acao")
    if acao == "alternar":
        execute("UPDATE plans SET is_active = 1 - is_active WHERE id=?", (pid,))
    elif acao == "preco":
        reais = request.form.get("price") or "0"
        try:
            cents = int(round(float(reais.replace(",", ".")) * 100))
            if cents >= 0:
                execute("UPDATE plans SET price_cents=? WHERE id=?", (cents, pid))
        except ValueError:
            pass
    elif acao == "excluir":
        if not query("SELECT id FROM promotions WHERE plan_id=?", (pid,), one=True):
            execute("DELETE FROM plans WHERE id=?", (pid,))
    audit(g.user["id"], f"plan_{acao}", "plan", pid, "", request.remote_addr)
    return redirect(url_for("admin.plans_admin"))

@bp.route("/admin/pagamentos")
def payments_admin():
    page = request.args.get("page", 1, type=int)
    per = 25
    total = query("SELECT COUNT(*) n FROM payments", one=True)["n"]
    rows = query("SELECT pay.*, u.nickname, tg.name AS group_name, pl.name AS plan_name "
                 "FROM payments pay LEFT JOIN users u ON u.id=pay.user_id "
                 "LEFT JOIN telegram_groups tg ON tg.id=pay.group_id "
                 "LEFT JOIN plans pl ON pl.id=pay.plan_id "
                 "ORDER BY pay.created_at DESC LIMIT ? OFFSET ?", (per, (page-1)*per))
    return render_template("admin/pagamentos.html", rows=rows, page=page,
                           pages=math.ceil(total/per) if total else 0)


@bp.route("/admin/destaques")
def destaques_admin():
    rows = query(
        "SELECT pr.id, pr.group_id, pr.payment_id, pr.plan_id, "
        "pr.starts_at, pr.ends_at, pr.is_active, "
        "tg.name AS group_name, tg.link AS group_link, "
        "u.nickname, pay.amount_cents, pay.status AS payment_status, "
        "pay.provider, pay.paid_at, pl.name AS plan_name "
        "FROM promotions pr "
        "LEFT JOIN telegram_groups tg ON tg.id = pr.group_id "
        "LEFT JOIN payments pay ON pay.id = pr.payment_id "
        "LEFT JOIN users u ON u.id = pay.user_id "
        "LEFT JOIN plans pl ON pl.id = pr.plan_id "
        "WHERE pr.is_active = 1 AND pr.ends_at > ? "
        "ORDER BY pr.ends_at ASC", (now(),))
    return render_template("admin/destaques.html", rows=rows)

@bp.route("/admin/destaques/<int:pid>/remover", methods=["POST"])
def destaque_remover(pid):
    if not csrf_ok():
        return "CSRF inválido", 400
    row = query(
        "SELECT id, group_id, is_active FROM promotions WHERE id = ?",
        (pid,), one=True)
    if not row:
        abort(404)
    if row["is_active"]:
        execute("UPDATE promotions SET is_active = 0 WHERE id = ?", (pid,))
        audit(g.user["id"], "remover_destaque", "promotion", pid,
              "Destaque removido pelo painel administrativo",
              request.remote_addr)
        flash("Destaque removido. O grupo e o histórico de pagamento foram preservados.", "ok")
    else:
        flash("Este destaque já estava desativado.", "ok")
    return redirect(url_for("admin.destaques_admin"))

@bp.route("/admin/banners", methods=["GET", "POST"])
def banners_admin():
    if request.method == "POST":
        if not csrf_ok():
            return "CSRF inválido", 400
        title = (request.form.get("title") or "").strip()[:80]
        link = (request.form.get("link") or "").strip()[:300]
        position = request.form.get("position") or "home_top"
        image = None
        if request.files.get("image"):
            image = save_image(request.files["image"], "banner")
        if title and link.startswith(("http://", "https://")) and image:
            execute("INSERT INTO banners(title, image_path, link, position, created_at) VALUES (?,?,?,?,?)",
                    (title, image, link, position, now()))
            flash("Banner criado.", "ok")
        else:
            flash("Banner precisa de título, link válido (http/https) e imagem.", "erro")
        return redirect(url_for("admin.banners_admin"))
    rows = query("SELECT * FROM banners ORDER BY created_at DESC")
    return render_template("admin/banners.html", rows=rows)

@bp.route("/admin/banners/<int:bid>/acao", methods=["POST"])
def banner_action(bid):
    if not csrf_ok():
        return "CSRF inválido", 400
    acao = request.form.get("acao")
    if acao == "alternar":
        execute("UPDATE banners SET is_active = 1 - is_active WHERE id=?", (bid,))
    elif acao == "excluir":
        execute("DELETE FROM banners WHERE id=?", (bid,))
    audit(g.user["id"], f"banner_{acao}", "banner", bid, "", request.remote_addr)
    return redirect(url_for("admin.banners_admin"))

@bp.route("/admin/config", methods=["GET", "POST"])
def settings_admin():
    if request.method == "POST":
        if not csrf_ok():
            return "CSRF inválido", 400
        for key in ("site_name", "site_tagline", "contact_email"):
            val = (request.form.get(key) or "").strip()[:200]
            execute("INSERT INTO settings(key, value) VALUES (?,?) "
                    "ON CONFLICT(key) DO UPDATE SET value = excluded.value", (key, val))
        flash("Configurações salvas.", "ok")
        return redirect(url_for("admin.settings_admin"))
    vals = {r["key"]: r["value"] for r in query("SELECT key, value FROM settings")}
    return render_template("admin/config.html", vals=vals)

@bp.route("/admin/logs")
def logs_admin():
    page = request.args.get("page", 1, type=int)
    per = 50
    total = query("SELECT COUNT(*) n FROM audit_log", one=True)["n"]
    rows = query("SELECT a.*, u.nickname FROM audit_log a LEFT JOIN users u ON u.id=a.actor_id "
                 "ORDER BY a.id DESC LIMIT ? OFFSET ?", (per, (page-1)*per))
    return render_template("admin/logs.html", rows=rows, page=page,
                           pages=math.ceil(total/per) if total else 0)

@bp.route("/admin/contatos")
def contacts_admin():
    rows = query("SELECT * FROM contact_messages ORDER BY created_at DESC LIMIT 100")
    return render_template("admin/contatos.html", rows=rows)

@bp.route("/admin/contatos/<int:cid>/lido", methods=["POST"])
def contact_read(cid):
    if csrf_ok():
        execute("UPDATE contact_messages SET handled=1 WHERE id=?", (cid,))
    return redirect(url_for("admin.contacts_admin"))
